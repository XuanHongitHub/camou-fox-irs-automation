from __future__ import annotations

import csv
import json
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, Iterable, List

from .logging_utils import ensure_dir

RESULT_HEADERS = [
    "batch_id",
    "record_id",
    "record_name",
    "status",
    "attempt_count",
    "proxy_used",
    "proxy_ip",
    "started_at",
    "ended_at",
    "last_step",
    "artifact_dir",
    "error_type",
    "error_code",
    "error_message",
    "confirmation_number",
    "step6_ein",
    "step6_legal_name",
    "step6_name_control",
    "step6_phone_number",
    "step6_county",
    "step6_state",
    "step6_start_date",
    "step6_principal_activity",
    "step6_principal_product_service",
    "step6_reason_for_applying",
    "step6_physical_location",
    "step6_responsible_name",
    "step6_responsible_ssn_itin",
    "step6_data_json",
    "pdf_path",
    "final_pdf_path",
]


def append_result_row(csv_path: Path, row: Dict[str, Any]) -> None:
    ensure_dir(csv_path.parent)
    write_header = not csv_path.exists()
    with csv_path.open("a", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=RESULT_HEADERS)
        if write_header:
            writer.writeheader()
        writer.writerow({k: row.get(k, "") for k in RESULT_HEADERS})
    _sync_styled_xlsx(csv_path)


def _sync_styled_xlsx(csv_path: Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill
    except Exception:
        return

    if not csv_path.exists():
        return

    with csv_path.open("r", newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        rows = list(reader)

    wb = Workbook()
    ws = wb.active
    ws.title = "results"
    ws.append(RESULT_HEADERS)

    header_fill = PatternFill(fill_type="solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    for c in range(1, len(RESULT_HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = header_fill
        cell.font = header_font

    green_fill = PatternFill(fill_type="solid", fgColor="DCFCE7")
    red_fill = PatternFill(fill_type="solid", fgColor="FEE2E2")
    yellow_fill = PatternFill(fill_type="solid", fgColor="FEF9C3")

    for idx, row in enumerate(rows, start=2):
        values = [row.get(k, "") for k in RESULT_HEADERS]
        ws.append(values)
        status = (row.get("status") or "").strip().lower()
        if status in {"success", "done"}:
            row_fill = green_fill
        elif status in {"manual_required", "cancelled"}:
            row_fill = yellow_fill
        else:
            row_fill = red_fill
        for c in range(1, len(RESULT_HEADERS) + 1):
            ws.cell(row=idx, column=c).fill = row_fill

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(str(csv_path.with_suffix(".xlsx")))


def write_json(path: Path, payload: Dict[str, Any]) -> None:
    ensure_dir(path.parent)
    path.write_text(json.dumps(payload, ensure_ascii=True, indent=2), encoding="utf-8")


def write_jsonl(path: Path, rows: Iterable[Dict[str, Any]]) -> None:
    ensure_dir(path.parent)
    with path.open("a", encoding="utf-8") as f:
        for row in rows:
            f.write(json.dumps(row, ensure_ascii=True) + "\n")


def load_json(path: Path) -> Dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def append_batch_state(state_path: Path, row: Dict[str, Any]) -> None:
    state: List[Dict[str, Any]] = []
    if state_path.exists():
        payload = json.loads(state_path.read_text(encoding="utf-8"))
        state = payload.get("rows", [])
    state.append(row)
    write_json(state_path, {"rows": state})


def update_batch_manifest(output_csv_path: Path, row: Dict[str, Any]) -> Path:
    """
    Keep a lightweight per-batch manifest for operators.
    """
    batch_id = str(row.get("batch_id", "")).strip()
    if not batch_id:
        raise ValueError("batch_id missing for manifest")
    out_dir = output_csv_path.parent / "reports"
    ensure_dir(out_dir)
    path = out_dir / f"{batch_id}_manifest.json"

    payload: Dict[str, Any] = {}
    if path.exists():
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            payload = {}

    total = int(payload.get("total", 0))
    success = int(payload.get("success", 0))
    failed = int(payload.get("failed", 0))
    status = str(row.get("status", "")).lower()
    total += 1
    if status in {"success", "done"}:
        success += 1
    else:
        failed += 1

    payload.update(
        {
            "batch_id": batch_id,
            "updated_at": datetime.now().isoformat(),
            "total": total,
            "success": success,
            "failed": failed,
            "last_record_id": str(row.get("record_id", "")),
            "last_status": status,
            "last_error_code": str(row.get("error_code", "")),
        }
    )
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path
